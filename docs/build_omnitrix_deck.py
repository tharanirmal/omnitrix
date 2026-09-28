"""Build the 8-slide Omnitrix deck (16:9) with python-pptx.

    pip install python-pptx
    python docs/build_omnitrix_deck.py            # writes docs/Omnitrix-final.pptx

Screenshots are read from "Claude outputs/" at the repo root; a missing file becomes a labelled placeholder.
Every number is traced in docs/deck-source.md.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "Claude outputs"
OUT = ROOT / "docs" / "Omnitrix-final.pptx"

# ------------------------------------------------------------------------------------------------ theme
BG, CARD, BORDER = "0B0F17", "161B22", "30363D"
TEXT, MUTED = "F0F6FC", "8B949E"
BLUE, MINT, CORAL = "38BDF8", "34D399", "F87171"       # electric cyan, mint emerald, coral red
MINT_TINT = "1A3130"                                     # mint at ~12% over the card fill
FONT = "Calibri"
W, H = 13.333, 7.5
M = 0.6                                   # side margin
CW = W - 2 * M                            # content width
FOOTER = "Omnitrix  ·  MSRIT Hackathon  ·  Track 1 Sovereign AI"


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ------------------------------------------------------------------------------------------------ helpers
def card(slide, x, y, w, h, fill=CARD, border=BORDER, border_pt=1.0, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if border:
        shp.line.color.rgb = rgb(border)
        shp.line.width = Pt(border_pt)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False            # no theme shadow
    return shp


def text(slide, x, y, w, h, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, italic=False):
    """paras: a string, or a list of paragraphs; a paragraph is a string or a list of
    (text, size, color, bold) runs (missing fields fall back to the defaults)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for run in [(para,)] if isinstance(para, str) else para:
            t, s, c, b = (*run, None, None, None)[:4]
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(s or size)
            r.font.bold = bold if b is None else b
            r.font.italic = italic
            r.font.color.rgb = rgb(c or color)
    return tb


def bullets(slide, x, y, w, h, items, dot=BLUE, size=15, color=TEXT, space=12):
    text(slide, x, y, w, h, [[("●   ", size - 5, dot, True), (t, size, color, False)] for t in items],
         space_after=space)


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    c.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))


def new_slide(prs, n, kicker=None, title=None, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])          # blank
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if kicker:
        text(s, M, 0.32, CW, 0.3, kicker, size=11, color=BLUE, bold=True)
    if title:
        text(s, M, 0.6, CW, 0.65, title, size=32, bold=True)
    if subtitle:
        text(s, M, 1.22, CW, 0.35, subtitle, size=14, color=MUTED)
    if n > 1:
        text(s, M, 7.02, 7.0, 0.25, FOOTER, size=9, color=MUTED)
        text(s, W - M - 1.0, 7.02, 1.0, 0.25, f"{n:02d}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def cell_border(cell, color=BORDER, width_pt=1.0):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, tag in enumerate(("lnL", "lnR", "lnT", "lnB")):       # borders precede the fill in tcPr
        tcPr.insert(i, parse_xml(
            f'<a:{tag} {nsdecls("a")} w="{int(width_pt * 12700)}"><a:solidFill><a:srgbClr val="{color}"/>'
            f'</a:solidFill></a:{tag}>'))


def cell_text(cell, value, size=12, color=TEXT, bold=False, fill=CARD):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.15)
    cell.margin_top = cell.margin_bottom = Inches(0.06)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ""
    r = p.add_run()
    r.text = value
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = rgb(color)
    cell_border(cell)


def screenshot(slide, path, x, y, w, h, keep=0.8):
    """The image inside a frame, cropped at the bottom to `keep` of its height, centred at its aspect."""
    card(slide, x, y, w, h, fill=BG, radius=0.04)
    pad = 0.08
    if path.exists():
        aspect = 1280 / (800 * keep)                           # the screenshots are 1280 x 800
        ih = h - 2 * pad
        iw = min(w - 2 * pad, ih * aspect)
        ih = iw / aspect
        pic = slide.shapes.add_picture(str(path), Inches(x + (w - iw) / 2), Inches(y + (h - ih) / 2),
                                       Inches(iw), Inches(ih))
        pic.crop_bottom = 1 - keep
    else:
        text(slide, x, y, w, h, f"Screenshot: {path.name}", size=12, color=MUTED, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------------------------------------------ slides
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)

# 1 · Title ----------------------------------------------------------------------------------------------
s = new_slide(prs, 1)
text(s, M, 2.0, CW, 0.3, "SOVEREIGN AI  ·  PERSONAL SECRETARY", size=12, color=BLUE, bold=True,
     align=PP_ALIGN.CENTER)
text(s, M, 2.35, CW, 1.2, "Omnitrix", size=72, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.5, 3.65, W - 3.0, 0.6, "A local AI secretary that remembers, plans, and acts — only with your OK.",
     size=22, color=BLUE, align=PP_ALIGN.CENTER)
badges = [("ENGINE", "engram — sovereign local brain", MINT),
          ("EVENT", "MSRIT Hackathon · Track 1: Sovereign AI", MINT),
          ("FOOTPRINT", "100% local · zero cloud calls", MINT)]
bw, bg_ = 3.6, 0.35
x0 = (W - (3 * bw + 2 * bg_)) / 2
for i, (label, value, accent) in enumerate(badges):
    x = x0 + i * (bw + bg_)
    card(s, x, 4.65, bw, 1.15, border=MINT)
    text(s, x + 0.25, 4.83, bw - 0.5, 0.25, label, size=11, color=accent, bold=True)
    text(s, x + 0.25, 5.1, bw - 0.5, 0.6, value, size=15, bold=True)

# 2 · Why Omnitrix -----------------------------------------------------------------------------------------
s = new_slide(prs, 2, "02  ·  WHY OMNITRIX", "Storing Data is Easy. Making Judgements is Hard.",
              "Why today's memory tools fail at long-term personal memory")
colw, gap = (CW - 0.33) / 2, 0.33
sides = [("TODAY'S MEMORY TOOLS", "Store first, judge never", CORAL,
          ["Mem0: 97.8% of 10,134 audited memories were junk",
           "Graphiti: a small-model judge invalidated 41% of ~3,950 facts, silently",
           "Local + cloud splits still leak private data on 7.5% of queries"]),
         ("OMNITRIX SOVEREIGN ARCHITECTURE", "Judge first, locally", MINT,
          ["Fine-tuned local System 1 (1.7B) makes the fast micro-judgements",
           "System 2 (14B) takes complex reasoning and drafting",
           "A human-in-the-loop gate before any action, with zero cloud calls"])]
for i, (label, heading, accent, items) in enumerate(sides):
    x = M + i * (colw + gap)
    card(s, x, 1.75, colw, 3.55, border=accent)
    text(s, x + 0.4, 2.05, colw - 0.8, 0.3, label, size=12, color=accent, bold=True)
    text(s, x + 0.4, 2.38, colw - 0.8, 0.5, heading, size=22, bold=True)
    bullets(s, x + 0.4, 3.15, colw - 0.8, 2.0, items, dot=accent, size=16)
card(s, M, 5.65, CW, 0.75, radius=0.5)
text(s, M, 5.65, CW, 0.75, "The hard part isn't storing things. It's the thousands of small judgements.",
     size=16, color=TEXT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)

# 3 · Architecture -------------------------------------------------------------------------------------------
s = new_slide(prs, 3, "03  ·  SYSTEM ARCHITECTURE", "One Laptop. One Brain. Seven Roles.")
nodes = [("INPUT", "Sources", "Email · Notes · Meetings", MUTED),
         ("S0", "Code filter", "Parse, dedupe, dates, rules", BLUE),
         ("S1", "Small judge", "Fine-tuned 1.7B, one token", BLUE),
         ("S2", "14B model", "The unsure middle, all writing", BLUE),
         ("H", "Your approval", "Herald watch or web app", CORAL),
         ("LEDGER", "Audit ledger", "Hash-chained, every decision", MINT)]
nw, ng, ny, nh = 1.75, (CW - 6 * 1.75) / 5, 2.25, 1.45
bx = M + nw + ng - 0.12                                        # engram brackets S0–S2
card(s, bx, 1.82, 3 * nw + 2 * ng + 0.24, 2.05, fill=None, border=BLUE, border_pt=1.5, radius=0.05)
text(s, bx + 0.15, 1.9, 3.0, 0.25, "ENGRAM  ·  THE BRAIN", size=10, color=BLUE, bold=True)
for i, (tag, name, detail, accent) in enumerate(nodes):
    x = M + i * (nw + ng)
    card(s, x, ny, nw, nh)
    text(s, x + 0.18, ny + 0.17, nw - 0.36, 0.22, tag, size=10, color=accent, bold=True)
    text(s, x + 0.18, ny + 0.42, nw - 0.36, 0.32, name, size=15, bold=True)
    text(s, x + 0.18, ny + 0.8, nw - 0.36, 0.55, detail, size=11, color=MUTED)
    if i < len(nodes) - 1:
        arrow(s, x + nw + 0.04, ny + nh / 2, x + nw + ng - 0.04, ny + nh / 2)
text(s, M, 4.3, CW, 0.3, "7 FUNCTIONAL ROLES ON ONE PIPELINE", size=11, color=MUTED, bold=True)
roles = [("Librarian", "Ingests and files everything"), ("Researcher", "Answers with cited sources"),
         ("Memory", "Keeps promises and decisions"), ("Planner", "Plans your day with a solver"),
         ("Operator", "Acts in your tools"), ("Guardian", "Gates every action"),
         ("Herald", "Reaches you on the watch")]
rg = 0.2
rw = (CW - 6 * rg) / 7
for i, (name, desc) in enumerate(roles):
    x = M + i * (rw + rg)
    card(s, x, 4.68, rw, 1.45)
    text(s, x + 0.15, 4.88, rw - 0.3, 0.32, name, size=15, color=BLUE, bold=True)
    text(s, x + 0.15, 5.28, rw - 0.3, 0.75, desc, size=12, color=MUTED)

# 4 · Efficiency ----------------------------------------------------------------------------------------------
s = new_slide(prs, 4, "04  ·  LOCAL MODEL PERFORMANCE", "A 1.7B Fine-Tuned Model Beats a 14B General Model")
metrics = [("54 min", "Local LoRA training", "on the Mac with MLX, peak 18 GB", BLUE),
           ("5,958", "Behavioural labels", "from the owner's own replies and filing, and EnronQA answers", MINT),
           ("3–5×", "Faster than the 14B", "zero-shot, on the same decisions", BLUE)]
mg = 0.4
mw = (CW - 2 * mg) / 3
for i, (num, label, desc, accent) in enumerate(metrics):
    x = M + i * (mw + mg)
    card(s, x, 1.75, mw, 2.45)
    text(s, x + 0.35, 1.98, mw - 0.7, 0.9, num, size=44, color=accent, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.35, 2.98, mw - 0.7, 0.35, label, size=16, bold=True)
    text(s, x + 0.35, 3.38, mw - 0.7, 0.7, desc, size=12, color=MUTED)
card(s, M, 4.45, CW, 2.3)
text(s, M + 0.4, 4.65, CW - 0.8, 0.28, "WHY A SMALL LOCAL JUDGE WINS", size=11, color=MINT, bold=True)
text(s, M + 0.4, 4.98, CW - 0.8, 1.7,
     [[("We fine-tune skills, never facts: knowledge stays in local Postgres, where it can be cited and deleted.",
        15, TEXT, True)],
      [("Trained on four decisions, the fine-tuned 1.7B beats the zero-shot 14B on every one: supported? "
        "96.6% vs 95.5% · relevant? 90.9% vs 87.5% · will reply? 60.1% vs 55.4% · folder 52.7% vs 47.9%.",
        15, TEXT, False)],
      [("Fused and quantized to 4-bit it is a 934 MB model, 2.9× faster again with no loss in accuracy.",
        14, MUTED, False)]], space_after=8)

# 5 · Verification & gated safety --------------------------------------------------------------------------------
s = new_slide(prs, 5, "05  ·  VERIFICATION & GATED SAFETY", "Checked Answers. Gated Actions.")
halves = [("ANSWERS CHECKED", "84% vs 6%", "right when the small judge vouches vs when it doesn't", BLUE,
           ["AUROC 0.92 on 100 held-out questions",
            "The local judge checks each answer against the emails it cites"]),
          ("ACTIONS GATED", "40.0% → 6.7%", "harmful side effects on 60 fresh sandbox tasks", MINT,
           ["0.62 owner prompts per task (approve-all: 0.97)",
            "Code sets the risk · the 14B checks intent · you approve",
            "Herald on Wear OS: approved on the wrist, ledgered 19 s after the gate paused"])]
for i, (label, big, sub, accent, items) in enumerate(halves):
    x = M + i * (colw + gap)
    card(s, x, 1.75, colw, 3.75, border=accent)
    text(s, x + 0.4, 2.05, colw - 0.8, 0.28, label, size=12, color=accent, bold=True)
    text(s, x + 0.4, 2.38, colw - 0.8, 0.9, big, size=40, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.4, 3.32, colw - 0.8, 0.5, sub, size=15, color=MUTED)
    bullets(s, x + 0.4, 4.0, colw - 0.8, 1.4, items, dot=accent, size=14, space=8)
card(s, M, 5.8, CW, 0.7, radius=0.5)
text(s, M, 5.8, CW, 0.7, "A small model may raise suspicion, never grant permission.", size=16,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)

# 6 · Interface showcase ---------------------------------------------------------------------------------------
s = new_slide(prs, 6, "06  ·  WHAT YOU SEE", "Interface Showcase",
              "Local web dashboard, constellation map and cryptographic ledger")
frames = [("1280-dark-06-answer.png", "Constellation Graph (3D Memory Map)"),
          ("1280-dark-07-atlas.png", "Schema Atlas & SQL Inspector"),
          ("1280-dark-09-agents.png", "Bitemporal Belief Timeline"),        # this file shows the belief timeline
          ("1280-dark-16-audit-verified.png", "Tamper-Evident Cryptographic Audit Log")]
fw, fh, fg = (CW - 0.33) / 2, 2.2, 0.33
for i, (fname, caption) in enumerate(frames):
    col, row = i % 2, i // 2
    x, y = M + col * (fw + fg), 1.72 + row * (fh + 0.45)
    screenshot(s, SHOTS / fname, x, y, fw, fh)
    text(s, x, y + fh + 0.08, fw, 0.28, caption, size=12, bold=True, align=PP_ALIGN.CENTER)

# 7 · Sovereign & comparison -------------------------------------------------------------------------------------
s = new_slide(prs, 7, "07  ·  SOVEREIGN BY DESIGN", "Sovereign by Design. Ready for DPDP.",
              "India's DPDP Rules: 1-year access logs · 72-hour breach reports · penalties up to ₹250 crore")
specs = ["Postgres + pgvector local storage", "Ollama + MLX local runtime", "Direct Wi-Fi Wear OS connection",
         "Hash-chained audit ledger"]
sg = 0.25
sw = (CW - 3 * sg) / 4
for i, spec in enumerate(specs):
    x = M + i * (sw + sg)
    card(s, x, 1.75, sw, 1.1)
    text(s, x + 0.25, 1.9, sw - 0.5, 0.25, f"0{i + 1}", size=12, color=BLUE, bold=True)
    text(s, x + 0.25, 2.17, sw - 0.5, 0.6, spec, size=14, bold=True)
rows = [("Tool", "Methodology", "Known vulnerability / limitation", "Omnitrix advantage"),
        ("Mem0", "ADD-only memory; local server sunset (Jul 2026)", "97.8% junk in a 10,134-memory audit",
         "Small-judge write gate: 151 → 36 beliefs on 60 emails"),
        ("Graphiti", "A small-model judge updates facts", "Invalidated 41% of ~3,950 facts, silently",
         "Contradictions are shown to you, never silently retired"),
        ("GBrain", "Postgres memory; LLM features hosted by default", "No per-decision confidence",
         "A calibrated probability on every decision, in the ledger"),
        ("Omnitrix", "Fine-tuned local judge + 14B + hash-chained ledger",
         "Measured weak spot: the “worth remembering?” gate is 72%", "All local; every decision auditable")]
gf = s.shapes.add_table(len(rows), 4, Inches(M), Inches(3.1), Inches(CW), Inches(3.3))
tbl = gf.table
tbl.first_row, tbl.horz_banding = True, False
for j, wdt in enumerate((1.55, 3.0, 3.65, CW - 1.55 - 3.0 - 3.65)):
    tbl.columns[j].width = Inches(wdt)
tbl.rows[0].height = Inches(0.45)
for r in range(1, len(rows)):
    tbl.rows[r].height = Inches(0.71)
for r, row in enumerate(rows):
    for c, value in enumerate(row):
        cell = tbl.cell(r, c)
        if r == 0:
            cell_text(cell, value.upper(), size=10, color=BLUE, bold=True, fill=CARD)
            continue
        ours = row[0] == "Omnitrix"
        fill = MINT_TINT if ours else BG
        if c == 0:
            cell_text(cell, value, size=13, color=MINT if ours else TEXT, bold=True, fill=fill)
        else:
            cell_text(cell, value, size=12, color=TEXT if ours else (MINT if c == 3 else MUTED), fill=fill)

# 8 · Demo, limits, next --------------------------------------------------------------------------------------
s = new_slide(prs, 8, "08  ·  DEMO, LIMITS, NEXT", "Five Minutes Live. Every Limit Measured.")
columns = [("DEMO · 5 MINUTES", BLUE, True, 5.0,
            ["Sign in to the kaminski brain",
             "Ask: “Where and at what time did Vince suggest meeting Michael Garberding on Tuesday?”",
             "Add an email with a promise → it becomes a belief",
             "Act: “Forward my most recent email from fatima to kofi” → the watch buzzes → approve",
             "Audit: the newest ledger row is owner:watch → Verify chain"]),
           ("LIMITS (MEASURED)", CORAL, False, 3.4,
            ["“Will they reply?” is 60.1%, so it stays a suggestion",
             "The “worth remembering?” gate: 72% (AUROC 0.80)",
             "The watch uses plain HTTP on the LAN",
             "Bluetooth on → the watch drops Wi-Fi when asleep"]),
           ("NEXT", MINT, False, CW - 5.0 - 3.4 - 0.54,
            ["Adapter v2: + remember, meeting_request, fulfilled, contradicts",
             "A learned “buzz now or later?” from your taps",
             "TLS pinned at pairing"])]
x = M
for label, accent, numbered, wdt, items in columns:
    card(s, x, 1.65, wdt, 3.7, border=accent)
    text(s, x + 0.3, 1.9, wdt - 0.6, 0.28, label, size=11, color=accent, bold=True)
    marks = [f"{k + 1}   " for k in range(len(items))] if numbered else ["●   "] * len(items)
    text(s, x + 0.3, 2.32, wdt - 0.6, 2.9,
         [[(m, 13 if numbered else 9, accent, True), (t, 13, TEXT, False)] for m, t in zip(marks, items)],
         space_after=9)
    x += wdt + 0.27
card(s, M, 5.75, CW, 0.6, radius=0.5)
text(s, M, 5.75, CW, 0.6, [[("Built this week:  ", 12, MUTED, True),
                            ("7,414 lines Python · 3,047 front end · 985 Kotlin · 15 SQL migrations · 113 tests",
                             12, TEXT, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print(f"wrote {OUT}")
